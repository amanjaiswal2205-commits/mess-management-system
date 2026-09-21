import os
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'messmgt.settings')
import django
django.setup()

import logging
from decimal import Decimal

logger = logging.getLogger(__name__)

import secrets
from datetime import timedelta

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth import get_user_model, login
from django.contrib.auth.models import User
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.core.mail import send_mail, get_connection
from django.db.models import Sum, Q
from django.http import HttpResponse, JsonResponse, HttpResponseRedirect
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.utils import timezone
from core.models import (
    UserProfile, MessSetting, AccountingPeriod, PeriodDefaultFee,
    Student, StudentPeriodAccount, Payment, Supplier, StockItem,
    Purchase, IssueToKitchen, Labour, LabourPayment, OtherExpense,
    MonthlySettlement, EmailOTP,
)
from core.decorators import (
    admin_required, viewer_or_admin_required, staff_or_admin_required,
    staff_permission_required, user_has_permission,
    frontend_management_restricted,
)
from core.forms import StudentForm, PaymentForm, PeriodDefaultFeeForm, StudentSearchForm, PurchaseForm, SignupForm, PeriodFilterForm, MessSettingForm, StudentPeriodAccountForm, DayWisePurchaseReportForm, RegistrationForm
from django.contrib import messages
from allauth.socialaccount.providers.google.views import GoogleOAuth2Adapter
from allauth.socialaccount.helpers import complete_social_login

PER_PAGE = 25

UserModel = get_user_model()


# Pagination helpers
def _build_query_string(request, exclude=None):
    exclude = exclude or []
    params = []
    for key in request.GET:
        if key in exclude:
            continue
        values = request.GET.getlist(key)
        for value in values:
            params.append(f"{key}={value}")
    if params:
        return "&".join(params) + "&"
    return ""


def _page_range_list(page_obj):
    paginator = page_obj.paginator
    current = page_obj.number
    total = paginator.num_pages
    if total <= 7:
        return list(range(1, total + 1))
    pages = []
    if current <= 3:
        pages.extend(range(1, 6))
        pages.append(None)
        pages.append(total)
    elif current >= total - 2:
        pages.append(1)
        pages.append(None)
        pages.extend(range(total - 4, total + 1))
    else:
        pages.append(1)
        pages.append(None)
        pages.extend(range(current - 1, current + 2))
        pages.append(None)
        pages.append(total)
    return pages


def _paginate(request, queryset, per_page=None, page_param='page'):
    from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
    if per_page is None:
        per_page = PER_PAGE
    paginator = Paginator(queryset, per_page)
    page = request.GET.get(page_param) or request.GET.get('p') or 1
    try:
        page_obj = paginator.page(page)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)
    page_range_list = _page_range_list(page_obj)
    query_string = _build_query_string(request, exclude=[page_param, 'p'])
    return {
        'page_obj': page_obj,
        'is_paginated': page_obj.has_other_pages(),
        'page_range_list': page_range_list,
        'query_string': query_string,
        'page_param': page_param,
    }


# Auth / misc
def signup(request):
    if request.method == 'POST':
        form = SignupForm(request.POST)
        if form.is_valid():
            form.save(request)
            messages.success(request, 'Account created successfully! Please login to continue.')
            return redirect('account_login')
    else:
        form = SignupForm()
    return render(request, 'core/signup.html', {'form': form})


def google_login(request):
    adapter = GoogleOAuth2Adapter(request)
    return complete_social_login(request, adapter)


def send_email_via_resend(to_email, subject, html_content, text_content=None):
    api_key = getattr(settings, 'RESEND_API_KEY', '').strip()
    from_email = getattr(settings, 'DEFAULT_FROM_EMAIL', '')
    timeout = getattr(settings, 'RESEND_API_TIMEOUT', 15)

    if not api_key:
        logger.error("RESEND_API_KEY is not configured")
        return False

    sender = from_email if from_email else 'onboarding@resend.dev'

    payload = {
        'from': sender,
        'to': [to_email],
        'subject': subject,
        'html': html_content,
    }
    if text_content:
        payload['text'] = text_content

    headers = {
        'Authorization': f'Bearer {api_key}',
        'Content-Type': 'application/json',
    }

    try:
        response = requests.post(
            'https://api.resend.com/emails',
            headers=headers,
            json=payload,
            timeout=timeout,
        )
        if response.status_code != 200:
            logger.error(
                "Resend API returned %s: %s",
                response.status_code,
                response.text,
            )
        response.raise_for_status()
        return True
    except Exception:
        logger.exception("Failed to send email via Resend API to %s", to_email)
        return False


def send_email_via_brevo(recipient_email, subject, html_content):
    api_key = getattr(settings, 'BREVO_API_KEY', '').strip()
    if not api_key:
        logger.error("BREVO_API_KEY is not configured")
        return False

    headers = {
        "accept": "application/json",
        "api-key": api_key,
        "content-type": "application/json",
    }

    payload = {
        "sender": {
            "name": "Mess Management System",
            "email": getattr(settings, 'DEFAULT_FROM_EMAIL', ''),
        },
        "to": [
            {
                "email": recipient_email
            }
        ],
        "subject": subject,
        "htmlContent": html_content,
    }

    try:
        response = requests.post(
            'https://api.brevo.com/v3/smtp/email',
            headers=headers,
            json=payload,
            timeout=getattr(settings, 'BREVO_API_TIMEOUT', 15),
        )
        if response.status_code not in (200, 201):
            logger.error(
                "Brevo API returned %s: %s",
                response.status_code,
                response.text,
            )
            return False
        return True
    except Exception:
        logger.exception("Failed to send email via Brevo API to %s", recipient_email)
        return False


def send_otp_email(email, otp, purpose, name=None):
    """Send a 6-digit OTP to `email` using the Brevo Email API.

    `purpose` is 'signup' or 'forgot_password' and controls the subject/body.
    `name` is the recipient's full name (used in signup emails).

    Returns True on success, False on failure.
    """
    if purpose == 'signup':
        subject = 'Verify Your Email | Hostel Mess Management System'
        greeting_name = name or email.split('@')[0]
        intro = (
            'Thank you for registering with the Hostel Mess Management System. '
            'Please use the One-Time Password (OTP) below to verify your email '
            'address and create your account:'
        )
        plain_body = (
            f"Dear {greeting_name},\n\n"
            f"{intro}\n\n"
            f"Your OTP is: {otp}\n\n"
            f"This OTP is valid for 10 minutes. Please do not share it with anyone.\n\n"
            f"Registered Gmail address: {email}\n\n"
            f"If you did not request this registration, you can safely ignore this email. "
            f"No changes will be made to your account.\n\n"
            f"Authorized Representative\n"
            f"Mess Committee\n"
            f"APJ Abdul Kalam Boys Hostel\n"
            f"https://hostelmess.in\n"
        )
        html_body = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Verify Your Email</title>
</head>
<body style="margin:0; padding:0; background-color:#f4f4f4; font-family:Arial, Helvetica, sans-serif;">
    <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="background-color:#f4f4f4;">
        <tr>
            <td align="center" style="padding:20px 10px;">
                <table role="presentation" width="600" cellpadding="0" cellspacing="0" border="0" style="background-color:#ffffff; border-radius:8px; overflow:hidden; max-width:600px;">
                    <tr>
                        <td style="background-color:#1a5276; padding:30px 20px; text-align:center;">
                            <h1 style="color:#ffffff; margin:0; font-size:22px; font-weight:bold; line-height:1.3;">APJ Abdul Kalam Boys Hostel</h1>
                            <p style="color:#d6eaf8; margin:8px 0 0 0; font-size:14px; font-weight:bold;">MESS COMMITTEE</p>
                            <p style="color:#aed6f1; margin:4px 0 0 0; font-size:12px;">Hostel Mess Management System</p>
                        </td>
                    </tr>
                    <tr>
                        <td style="padding:30px 20px;">
                            <p style="color:#333333; font-size:16px; margin:0 0 16px 0; line-height:1.5;">Dear {greeting_name},</p>
                            <p style="color:#555555; font-size:14px; line-height:1.6; margin:0 0 20px 0;">
                                {intro}
                            </p>
                            <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="margin:20px 0;">
                                <tr>
                                    <td align="center" style="background-color:#eaf2f8; border:2px dashed #1a5276; border-radius:8px; padding:24px;">
                                        <span style="font-size:36px; font-weight:bold; color:#1a5276; letter-spacing:8px;">{otp}</span>
                                    </td>
                                </tr>
                            </table>
                            <p style="color:#555555; font-size:14px; line-height:1.6; margin:0 0 16px 0;">
                                <strong style="color:#1a5276;">This OTP is valid for 10 minutes.</strong>
                            </p>
                            <p style="color:#555555; font-size:14px; line-height:1.6; margin:0 0 16px 0;">
                                Registered Gmail address: <strong>{email}</strong>
                            </p>
                            <p style="color:#c0392b; font-size:14px; line-height:1.6; margin:0 0 20px 0;">
                                <strong>Security Warning:</strong> Never share this OTP with anyone. 
                                Our team will never ask you for your OTP.
                            </p>
                            <p style="color:#777777; font-size:14px; line-height:1.6; margin:0;">
                                If you did not request this registration, please ignore this email. 
                                No changes will be made to your account.
                            </p>
                        </td>
                    </tr>
                    <tr>
                        <td style="background-color:#f4f4f4; padding:20px; text-align:center; border-top:1px solid #dddddd;">
                            <p style="color:#555555; font-size:12px; margin:0 0 8px 0; line-height:1.5;">
                                <strong>Authorized Representative</strong><br>
                                Mess Committee<br>
                                APJ Abdul Kalam Boys Hostel
                            </p>
                            <p style="color:#777777; font-size:12px; margin:0;">
                                <a href="https://hostelmess.in" style="color:#1a5276; text-decoration:none;">https://hostelmess.in</a>
                            </p>
                        </td>
                    </tr>
                </table>
            </td>
        </tr>
    </table>
</body>
</html>"""
    else:
        subject = 'Password Reset OTP | Hostel Mess Management System'
        greeting_name = name or email.split('@')[0]
        intro = (
            'We received a request to reset the password for your Hostel Mess '
            'Management System account. Please use the One-Time Password (OTP) '
            'below to continue:'
        )
        plain_body = (
            f"Dear {greeting_name},\n\n"
            f"{intro}\n\n"
            f"Your OTP is: {otp}\n\n"
            f"This OTP is valid for 10 minutes. Please do not share it with anyone.\n\n"
            f"Registered Gmail address: {email}\n\n"
            f"If you did not request this password reset, you can safely ignore this email. "
            f"No changes will be made to your account.\n\n"
            f"Authorized Representative\n"
            f"Mess Committee\n"
            f"APJ Abdul Kalam Boys Hostel\n"
            f"https://hostelmess.in\n"
        )
        html_body = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Password Reset OTP</title>
</head>
<body style="margin:0; padding:0; background-color:#f4f4f4; font-family:Arial, Helvetica, sans-serif;">
    <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="background-color:#f4f4f4;">
        <tr>
            <td align="center" style="padding:20px 10px;">
                <table role="presentation" width="600" cellpadding="0" cellspacing="0" border="0" style="background-color:#ffffff; border-radius:8px; overflow:hidden; max-width:600px;">
                    <tr>
                        <td style="background-color:#1a5276; padding:30px 20px; text-align:center;">
                            <h1 style="color:#ffffff; margin:0; font-size:22px; font-weight:bold; line-height:1.3;">APJ Abdul Kalam Boys Hostel</h1>
                            <p style="color:#d6eaf8; margin:8px 0 0 0; font-size:14px; font-weight:bold;">MESS COMMITTEE</p>
                            <p style="color:#aed6f1; margin:4px 0 0 0; font-size:12px;">Hostel Mess Management System</p>
                        </td>
                    </tr>
                    <tr>
                        <td style="padding:30px 20px;">
                            <p style="color:#333333; font-size:16px; margin:0 0 16px 0; line-height:1.5;">Dear {greeting_name},</p>
                            <p style="color:#555555; font-size:14px; line-height:1.6; margin:0 0 20px 0;">
                                {intro}
                            </p>
                            <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="margin:20px 0;">
                                <tr>
                                    <td align="center" style="background-color:#eaf2f8; border:2px dashed #1a5276; border-radius:8px; padding:24px;">
                                        <span style="font-size:36px; font-weight:bold; color:#1a5276; letter-spacing:8px;">{otp}</span>
                                    </td>
                                </tr>
                            </table>
                            <p style="color:#555555; font-size:14px; line-height:1.6; margin:0 0 16px 0;">
                                <strong style="color:#1a5276;">This OTP is valid for 10 minutes.</strong>
                            </p>
                            <p style="color:#555555; font-size:14px; line-height:1.6; margin:0 0 16px 0;">
                                Registered Gmail address: <strong>{email}</strong>
                            </p>
                            <p style="color:#c0392b; font-size:14px; line-height:1.6; margin:0 0 20px 0;">
                                <strong>Security Warning:</strong> Never share this OTP with anyone. 
                                Our team will never ask you for your OTP.
                            </p>
                            <p style="color:#777777; font-size:14px; line-height:1.6; margin:0;">
                                If you did not request this password reset, please ignore this email. 
                                No changes will be made to your account.
                            </p>
                        </td>
                    </tr>
                    <tr>
                        <td style="background-color:#f4f4f4; padding:20px; text-align:center; border-top:1px solid #dddddd;">
                            <p style="color:#555555; font-size:12px; margin:0 0 8px 0; line-height:1.5;">
                                <strong>Authorized Representative</strong><br>
                                Mess Committee<br>
                                APJ Abdul Kalam Boys Hostel
                            </p>
                            <p style="color:#777777; font-size:12px; margin:0;">
                                <a href="https://hostelmess.in" style="color:#1a5276; text-decoration:none;">https://hostelmess.in</a>
                            </p>
                        </td>
                    </tr>
                </table>
            </td>
        </tr>
    </table>
</body>
</html>"""

    return send_email_via_brevo(
        recipient_email=email,
        subject=subject,
        html_content=html_body,
    )


def register(request):
    if request.method == 'POST':
        form = RegistrationForm(request.POST)
        if form.is_valid():
            email = form.cleaned_data['email'].lower().strip()
            otp = f"{secrets.randbelow(1000000):06d}"

            EmailOTP.objects.filter(email=email, purpose='signup', is_verified=False).update(is_verified=True)

            expires_at = timezone.now() + timedelta(minutes=10)
            EmailOTP.objects.create(
                email=email,
                otp_hash=EmailOTP.hash_otp(otp),
                purpose='signup',
                expires_at=expires_at,
            )

            request.session['pending_registration'] = {
                'full_name': form.cleaned_data['full_name'],
                'email': email,
                'mobile': form.cleaned_data['mobile'],
                'password': form.cleaned_data['password1'],
            }

            try:
                send_otp_email(email, otp, 'signup', name=form.cleaned_data.get('full_name', ''))
            except Exception as exc:
                logger.error("Failed to send OTP email to %s: %s", email, exc)
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({'success': False, 'error': 'OTP email could not be sent. Please try again.'})
                messages.error(request, 'Failed to send OTP email. Please try again later.')
                return redirect('account_register')

            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({'success': True, 'otp_sent': True})
            return redirect('otp_verify_placeholder')
    else:
        form = RegistrationForm()
    return render(request, 'account/register.html', {'form': form})


def otp_verify_placeholder(request):
    pending = request.session.get('pending_registration')
    if not pending:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({'success': False, 'error': 'Your session has expired. Please register again.'})
        return redirect('account_register')

    email = pending.get('email', '').lower().strip() if pending.get('email') else ''
    otp_error = None

    if request.method == 'POST':
        submitted_otp = request.POST.get('otp', '').strip()
        is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'

        if not submitted_otp or len(submitted_otp) != 6 or not submitted_otp.isdigit():
            otp_error = 'Please enter a valid 6-digit OTP.'
        else:
            try:
                otp_record = EmailOTP.objects.filter(
                    email=email,
                    purpose='signup',
                    is_verified=False,
                ).latest('created_at')
            except EmailOTP.DoesNotExist:
                otp_record = None

            if not otp_record:
                otp_error = 'OTP has expired or already used. Please request a new one.'
            elif otp_record.expires_at < timezone.now():
                otp_error = 'OTP has expired. Please request a new OTP.'
            elif not otp_record.verify_otp(submitted_otp):
                otp_error = 'Wrong OTP. Please try again.'
            else:
                if UserModel.objects.filter(email=email).exists():
                    otp_error = 'An account already exists with this Gmail address. Please log in.'
                else:
                    full_name = pending.get('full_name', '')
                    password = pending.get('password', '')
                    mobile = pending.get('mobile', '')

                    user = UserModel.objects.create_user(
                        username=email,
                        email=email,
                        password=password,
                    )
                    user.first_name = full_name
                    user.last_name = ''
                    user.save()

                    profile, _ = UserProfile.objects.get_or_create(user=user)
                    profile.mobile = mobile
                    profile.save()

                    otp_record.is_verified = True
                    otp_record.save()

                    request.session.pop('pending_registration', None)

                    if is_ajax:
                        return JsonResponse({
                            'success': True,
                            'redirect_url': reverse('account_login'),
                            'message': 'Account created successfully. You can now log in.',
                        })
                    messages.success(request, 'Account created successfully. You can now log in.')
                    return redirect('account_login')

        if otp_error is not None and is_ajax:
            return JsonResponse({'success': False, 'error': otp_error})

    return render(request, 'account/otp_verify.html', {'email': email, 'otp_error': otp_error})


# ---------------------------------------------------------------------------
# Forgot password (email OTP based)
# ---------------------------------------------------------------------------

def _normalize_email(value):
    return (value or '').strip().lower()


def _validate_otp_format(otp):
    return bool(otp) and otp.isdigit() and len(otp) == 6
def _forgot_password_send_otp(request, email):
    """Create and email a fresh reset OTP. Any prior unused reset OTPs for this

    email are invalidated first so a code can't be reused and only one valid

    code exists at a time."""

    EmailOTP.objects.filter(
        email=email, purpose='forgot_password', is_verified=False
    ).update(is_verified=True)
    otp = f"{secrets.randbelow(1000000):06d}"
    expires_at = timezone.now() + timedelta(minutes=10)
    EmailOTP.objects.create(
        email=email,
        otp_hash=EmailOTP.hash_otp(otp),
        purpose='forgot_password',
        expires_at=expires_at,
    )
    try:
        user = UserModel.objects.get(email=email)
        name = user.get_full_name() or user.username
    except UserModel.DoesNotExist:
        name = None
    return send_otp_email(email, otp, 'forgot_password', name=name)

def forgot_password_request(request):
    if request.method == 'POST':
        email = _normalize_email(request.POST.get('email'))
        is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'

        if not email:
            error_msg = 'Please enter a valid email address.'
            if is_ajax:
                return JsonResponse({'success': False, 'error': error_msg})
            messages.error(request, error_msg)
            return redirect('forgot_password_request')

        try:
            from django.core.validators import validate_email as _validate_email
            _validate_email(email)
        except Exception:
            error_msg = 'Please enter a valid email address.'
            if is_ajax:
                return JsonResponse({'success': False, 'error': error_msg})
            messages.error(request, error_msg)
            return redirect('forgot_password_request')

        user = UserModel.objects.filter(email=email).first()
        if user is None:
            error_msg = 'No account found with this email address.'
            if is_ajax:
                return JsonResponse({'success': False, 'error': error_msg})
            messages.error(request, error_msg)
            return redirect('forgot_password_request')

        sent = _forgot_password_send_otp(request, email)
        if not sent:
            error_msg = 'Could not send the OTP email. Please try again later.'
            if is_ajax:
                return JsonResponse({'success': False, 'error': error_msg})
            messages.error(request, error_msg)
            return redirect('forgot_password_request')
        request.session['password_reset_email'] = email
        request.session.pop('password_reset_verified', None)

        if is_ajax:
            return JsonResponse(
                {'success': True, 'otp_sent': True, 'redirect_url': reverse('forgot_password_verify')}
            )
        return redirect('forgot_password_verify')

    return render(request, 'account/forgot_password_request.html', {})


def forgot_password_verify(request):
    email = request.session.get('password_reset_email')
    if not email:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({'success': False, 'error': 'Your session has expired. Please request a new OTP.'})
        messages.error(request, 'Your session has expired. Please request a new OTP.')
        return redirect('forgot_password_request')

    if request.method == 'POST':
        submitted_otp = (request.POST.get('otp') or '').strip()
        is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
        otp_error = None

        if not _validate_otp_format(submitted_otp):
            otp_error = 'Please enter a valid 6-digit OTP.'
        else:
            try:
                otp_record = EmailOTP.objects.filter(
                    email=email, purpose='forgot_password', is_verified=False
                ).latest('created_at')
            except EmailOTP.DoesNotExist:
                otp_record = None

            if not otp_record:
                otp_error = 'OTP has expired or already used. Please request a new OTP.'
            elif otp_record.expires_at < timezone.now():
                otp_error = 'OTP has expired. Please request a new OTP.'
            elif not otp_record.verify_otp(submitted_otp):
                otp_error = 'Wrong OTP. Please try again.'
            else:
                otp_record.is_verified = True
                otp_record.save()
                request.session['password_reset_verified'] = email
                request.session.pop('password_reset_email', None)
                if is_ajax:
                    return JsonResponse({'success': True, 'redirect_url': reverse('forgot_password_reset')})
                return redirect('forgot_password_reset')

        if otp_error and is_ajax:
            return JsonResponse({'success': False, 'error': otp_error})
        return render(request, 'account/forgot_password_verify.html', {'email': email, 'otp_error': otp_error})

    return render(request, 'account/forgot_password_verify.html', {'email': email})


def forgot_password_reset(request):
    email = request.session.get('password_reset_verified')
    if not email:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({'success': False, 'error': 'Your session has expired. Please request a new OTP.'})
        messages.error(request, 'Your session has expired. Please request a new OTP.')
        return redirect('forgot_password_request')

    if request.method == 'POST':
        password1 = request.POST.get('password1', '')
        password2 = request.POST.get('password2', '')
        is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
        errors = []

        if password1 != password2:
            errors.append('Passwords do not match.')
        try:
            validate_password(password1)
        except ValidationError as exc:
            errors.extend(exc.messages)

        if errors:
            if is_ajax:
                return JsonResponse({'success': False, 'error': ' '.join(errors)})
            return render(request, 'account/forgot_password_reset.html', {'email': email, 'errors': errors})

        user = UserModel.objects.filter(email=email).first()
        if user is None:
            if is_ajax:
                return JsonResponse({'success': False, 'error': 'Account not found. Please request a new OTP.'})
            messages.error(request, 'Account not found. Please request a new OTP.')
            return redirect('forgot_password_request')

        user.set_password(password1)
        user.save()

        # Invalidate every reset OTP for this email and drop the temporary
        # session authorization so the code/token cannot be reused.
        EmailOTP.objects.filter(email=email, purpose='forgot_password').update(is_verified=True)
        request.session.pop('password_reset_verified', None)
        request.session.pop('password_reset_email', None)

        messages.success(request, 'Password changed successfully. You can now log in.')
        if is_ajax:
            return JsonResponse({'success': True, 'redirect_url': reverse('account_login')})
        return redirect('account_login')

    return render(request, 'account/forgot_password_reset.html', {'email': email})


# Home & Dashboard
def home(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    return render(request, 'core/home.html')


@login_required
@viewer_or_admin_required
def dashboard(request):
    from django.db.models import Sum
    from django.db.models.functions import TruncMonth
    from collections import defaultdict
    import json

    profile = getattr(request.user, 'profile', None)

    total_fee = Payment.objects.aggregate(total=Sum('amount'))['total'] or 0

    total_expense = 0
    total_expense += Purchase.objects.aggregate(total=Sum('amount'))['total'] or 0
    total_expense += LabourPayment.objects.filter(status='PAID').aggregate(total=Sum('amount'))['total'] or 0
    total_expense += OtherExpense.objects.aggregate(total=Sum('amount'))['total'] or 0

    surplus_deficit = total_fee - total_expense

    fee_qs = Payment.objects.annotate(
        month_trunc=TruncMonth('month')
    ).values('month_trunc').annotate(total=Sum('amount')).order_by('month_trunc')

    purchase_expense = Purchase.objects.annotate(
        month_trunc=TruncMonth('bill_date')
    ).values('month_trunc').annotate(total=Sum('amount'))

    labour_expense = LabourPayment.objects.filter(status='PAID').annotate(
        month_trunc=TruncMonth('month')
    ).values('month_trunc').annotate(total=Sum('amount'))

    other_expense = OtherExpense.objects.annotate(
        month_trunc=TruncMonth('month')
    ).values('month_trunc').annotate(total=Sum('amount'))

    all_months = set()
    for entry in fee_qs:
        all_months.add(entry['month_trunc'])
    for entry in purchase_expense:
        all_months.add(entry['month_trunc'])
    for entry in labour_expense:
        all_months.add(entry['month_trunc'])
    for entry in other_expense:
        all_months.add(entry['month_trunc'])

    fee_lookup = {entry['month_trunc']: float(entry['total'] or 0) for entry in fee_qs}
    expense_lookup = defaultdict(float)
    for entry in purchase_expense:
        expense_lookup[entry['month_trunc']] += float(entry['total'] or 0)
    for entry in labour_expense:
        expense_lookup[entry['month_trunc']] += float(entry['total'] or 0)
    for entry in other_expense:
        expense_lookup[entry['month_trunc']] += float(entry['total'] or 0)

    sorted_months = sorted(all_months)
    months = [m.strftime('%b %Y') for m in sorted_months]
    fees = [fee_lookup.get(m, 0) for m in sorted_months]
    expenses = [expense_lookup.get(m, 0) for m in sorted_months]

    recent_purchases = Purchase.objects.select_related('supplier', 'item').all().order_by('-bill_date')[:5]

    context = {
        'user': request.user,
        'profile': profile,
        'total_fee': total_fee,
        'total_expense': total_expense,
        'surplus_deficit': surplus_deficit,
        'months': json.dumps(months),
        'fees': json.dumps(fees),
        'expenses': json.dumps(expenses),
        'recent_purchases': recent_purchases,
    }
    return render(request, 'core/dashboard.html', context)


# Students
@login_required
@viewer_or_admin_required
def student_list(request):
    qs = Student.objects.all().order_by('hostel_id')
    search = request.GET.get('search', '').strip()
    if search:
        qs = qs.filter(
            Q(hostel_id__icontains=search) |
            Q(student_name__icontains=search) |
            Q(room_no__icontains=search) |
            Q(user__first_name__icontains=search) |
            Q(user__last_name__icontains=search)
        )
    ctx = _paginate(request, qs, PER_PAGE)
    page_obj = ctx['page_obj']
    ctx['students'] = page_obj.object_list
    ctx['search_form'] = StudentSearchForm(initial={'query': search})
    ctx['search'] = search
    ctx['total_students'] = qs.count()
    ctx['named_students'] = qs.filter(Q(user__first_name__gt='') | Q(student_name__gt='')).count()
    ctx['unnamed_students'] = ctx['total_students'] - ctx['named_students']
    return render(request, 'core/student_list.html', ctx)


@admin_required
def student_import(request):
    """
    Import students from Excel file (.xlsx).
    Admin only.
    Expected Excel structure (MESS ST26.xlsx):
    - Sheet: Sheet1
    - Header row: Row 6
    - Columns B:G = Sr.No, Student ID, Student Name, Room no, Mobile no, Email ID
    - Data rows: Row 7 onwards (59 students)
    """
    from openpyxl import load_workbook
    from django.db import transaction
    from django.contrib.auth.models import User

    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']

        # Validate file type
        if not excel_file.name.endswith('.xlsx'):
            messages.error(request, 'Only .xlsx files are allowed.')
            return redirect('student_import')

        # Validate file size (max 5MB)
        if excel_file.size > 5 * 1024 * 1024:
            messages.error(request, 'File size must be less than 5MB.')
            return redirect('student_import')

        try:
            wb = load_workbook(excel_file, read_only=True)
            ws = wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active

            # Find header row by looking for "Student ID" in rows 1-10
            header_row = None
            for row_idx in range(1, 11):
                for cell in ws[row_idx]:
                    if cell.value and str(cell.value).strip().lower() == 'student id':
                        header_row = row_idx
                        break
                if header_row:
                    break

            if not header_row:
                messages.error(request, 'Could not find "Student ID" header in first 10 rows.')
                return redirect('student_import')

            # Get headers from detected header row
            headers = []
            for cell in ws[header_row]:
                if cell.value:
                    headers.append(str(cell.value).strip())
                else:
                    headers.append('')

            # Expected columns mapping by header name
            expected_columns = {
                'student id': 'student_id',
                'student name': 'student_name',
                'room no': 'room_no',
                'mobile no': 'mobile_no',
                'email id': 'email_id',
            }

            # Map header names to column indices
            col_map = {}
            for idx, header in enumerate(headers):
                header_lower = header.lower().strip()
                if header_lower in expected_columns:
                    col_map[expected_columns[header_lower]] = idx

            required_cols = ['student_id', 'student_name']
            missing_cols = [col for col in required_cols if col not in col_map]
            if missing_cols:
                messages.error(
                    request,
                    f'Missing required columns: {", ".join(missing_cols)}. '
                    f'Found headers: {", ".join([h for h in headers if h])}'
                )
                return redirect('student_import')

            # Statistics
            stats = {
                'total_rows': 0,
                'new_students': 0,
                'updated_students': 0,
                'skipped_rows': 0,
                'errors': []
            }

            # Process rows starting from row after header
            with transaction.atomic():
                for row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
                    stats['total_rows'] += 1

                    # Skip completely blank rows
                    if not any(cell for cell in row if cell is not None and str(cell).strip()):
                        stats['skipped_rows'] += 1
                        continue

                    # Use savepoint for each row to prevent one failure from breaking entire transaction
                    sid = transaction.savepoint()
                    try:
                        # Get values by column index, handle out-of-bounds
                        def get_cell(idx):
                            if idx < len(row) and row[idx] is not None:
                                val = str(row[idx]).strip()
                                return val if val else ''
                            return ''

                        student_id = get_cell(col_map.get('student_id', -1))
                        student_name = get_cell(col_map.get('student_name', -1))
                        room_no = get_cell(col_map.get('room_no', -1))
                        mobile_no = get_cell(col_map.get('mobile_no', -1))
                        email_id = get_cell(col_map.get('email_id', -1))

                        if not student_id or not student_name:
                            stats['skipped_rows'] += 1
                            stats['errors'].append(f'Row {row_idx}: Missing Student ID or Student Name')
                            transaction.savepoint_rollback(sid)
                            continue

                        # Clean mobile - remove non-digits, handle numbers stored as int/float
                        if mobile_no:
                            mobile_no = ''.join(filter(str.isdigit, str(mobile_no)))
                        # Clean email
                        if email_id:
                            email_id = email_id.strip()

                        # Convert empty strings to None for nullable fields
                        room_no = room_no if room_no else None
                        mobile_no = mobile_no if mobile_no else None
                        email_id = email_id if email_id else None

                        # Check if student exists by hostel_id
                        student = Student.objects.filter(hostel_id=student_id).first()

                        if student:
                            # Update existing student
                            student.student_name = student_name
                            if room_no:
                                student.room_no = room_no

                            # Preserve existing mobile/email if Excel cell is blank
                            # Only update when Excel provides a non-empty value
                            if mobile_no:
                                student.phone = mobile_no
                            if email_id:
                                student.email = email_id

                            student.save()

                            # Update linked user
                            user = student.user
                            parts = student_name.strip().split(' ', 1)
                            user.first_name = parts[0]
                            user.last_name = parts[1] if len(parts) > 1 else ''
                            if student.hostel_id:
                                user.username = student.hostel_id
                            user.save()

                            stats['updated_students'] += 1
                        else:
                            # Create new student with user
                            user = User.objects.create_user(
                                username=student_id,
                                email=email_id if email_id else '',
                                first_name=student_name.split(' ', 1)[0],
                                last_name=student_name.split(' ', 1)[1] if ' ' in student_name else '',
                            )
                            student = Student.objects.create(
                                user=user,
                                hostel_id=student_id,
                                student_name=student_name,
                                room_no=room_no,
                                phone=mobile_no,
                                email=email_id,
                                is_active=True,
                            )
                            stats['new_students'] += 1

                        transaction.savepoint_commit(sid)

                    except Exception as e:
                        transaction.savepoint_rollback(sid)
                        stats['skipped_rows'] += 1
                        stats['errors'].append(f'Row {row_idx}: {str(e)}')
                        continue

            # Prepare success message
            msg = (
                f'Import complete. '
                f'Total rows: {stats["total_rows"]}, '
                f'New students: {stats["new_students"]}, '
                f'Updated: {stats["updated_students"]}, '
                f'Skipped: {stats["skipped_rows"]}.'
            )
            if stats['errors']:
                msg += ' Errors: ' + '; '.join(stats['errors'][:5])
                if len(stats['errors']) > 5:
                    msg += f' ... and {len(stats["errors"]) - 5} more errors.'
            messages.success(request, msg)
            return redirect('student_import')

        except Exception as e:
            messages.error(request, f'Failed to process Excel file: {str(e)}')
            return redirect('student_import')

    return render(request, 'core/student_import.html')


@login_required
@staff_permission_required('can_export_collection_excel')
def students_export_excel(request):
    return HttpResponse("export students", content_type='text/plain')


@frontend_management_restricted
def student_add(request):
    if request.method == 'POST':
        form = StudentForm(request.POST)
        if form.is_valid():
            try:
                form.save()
                messages.success(request, 'Student added.')
                return redirect('student_list')
            except Exception as e:
                messages.error(request, f'Could not save student: {e}')
    else:
        form = StudentForm()
    return render(request, 'core/student_form.html', {'form': form})


@frontend_management_restricted
def student_edit(request, pk):
    student = get_object_or_404(Student, pk=pk)
    if request.method == 'POST':
        form = StudentForm(request.POST, instance=student, user_instance=student.user)
        if form.is_valid():
            form.save()
            messages.success(request, 'Student updated.')
            return redirect('student_list')
    else:
        form = StudentForm(instance=student, user_instance=student.user)
    return render(request, 'core/student_form.html', {'form': form, 'student': student})


@frontend_management_restricted
def student_toggle_active(request, pk):
    student = get_object_or_404(Student, pk=pk)
    student.is_active = not student.is_active
    student.save()
    return redirect('student_list')


@frontend_management_restricted
def mess_setting(request):
    setting = MessSetting.objects.first()
    form = MessSettingForm(instance=setting)
    if request.method == 'POST':
        form = MessSettingForm(request.POST, instance=setting)
        if form.is_valid():
            form.save()
            messages.success(request, 'Mess settings updated.')
            return redirect('student_list')
    return render(request, 'core/mess_setting_form.html', {'form': form, 'setting': setting})


# Period Default Fee
@frontend_management_restricted
def period_default_fee_set(request, period_id):
    period = get_object_or_404(AccountingPeriod, pk=period_id)
    obj, created = PeriodDefaultFee.objects.get_or_create(period=period)
    if request.method == 'POST':
        form = PeriodDefaultFeeForm(request.POST, instance=obj)
        if form.is_valid():
            form.save()
            messages.success(request, 'Default fee saved.')
            return redirect('dashboard')
    else:
        form = PeriodDefaultFeeForm(instance=obj)
    return render(request, 'core/period_default_fee_form.html', {'form': form, 'period': period})


@frontend_management_restricted
def period_default_fee_apply_unset(request, period_id):
    period = get_object_or_404(AccountingPeriod, pk=period_id)
    default_fee_obj = getattr(period, 'default_fee', None)
    if not default_fee_obj:
        messages.error(request, 'No default fee set for this period.')
        return redirect('dashboard')

    default_fee = default_fee_obj.default_fee_per_student
    active_students = Student.objects.filter(is_active=True)

    created_count = 0
    with transaction.atomic():
        for student in active_students:
            account, created = StudentPeriodAccount.objects.get_or_create(
                student=student,
                period=period,
                defaults={'total_to_collect': default_fee}
            )
            if created:
                created_count += 1

    messages.success(request, f'Default fee applied to {created_count} students without existing accounts.')
    return redirect('period_default_fee_set', period_id=period.pk)


@frontend_management_restricted
def period_default_fee_apply_all(request, period_id):
    period = get_object_or_404(AccountingPeriod, pk=period_id)
    default_fee_obj = getattr(period, 'default_fee', None)
    if not default_fee_obj:
        messages.error(request, 'No default fee set for this period.')
        return redirect('dashboard')

    default_fee = default_fee_obj.default_fee_per_student
    active_students = Student.objects.filter(is_active=True)

    updated_count = 0
    created_count = 0
    with transaction.atomic():
        for student in active_students:
            account, created = StudentPeriodAccount.objects.get_or_create(
                student=student,
                period=period,
                defaults={'total_to_collect': default_fee}
            )
            if not created:
                account.total_to_collect = default_fee
                account.is_manual_remaining = False
                account.manual_remaining = None
                account.save()
                updated_count += 1
            else:
                created_count += 1

    messages.success(request, f'Default fee applied: {created_count} new accounts created, {updated_count} existing accounts updated.')
    return redirect('period_default_fee_set', period_id=period.pk)


# Payments
@login_required
@viewer_or_admin_required
def payments_list(request):
    qs = Payment.objects.all().order_by('-created_at')
    ctx = _paginate(request, qs, PER_PAGE)
    ctx['payments'] = ctx['page_obj'].object_list
    return render(request, 'core/payments_list.html', ctx)


@login_required
@staff_permission_required('can_export_payment_excel')
def payments_export_excel(request):
    return HttpResponse("export payments", content_type='text/plain')


@login_required
@viewer_or_admin_required
def payment_summary(request):
    from django.db.models import Sum, F, DecimalField
    from django.db.models.functions import Coalesce

    period_form = PeriodFilterForm(request.GET)
    selected_period = None
    summary = []

    if period_form.is_valid():
        selected_period = period_form.cleaned_data.get('period')

    if selected_period:
        accounts = StudentPeriodAccount.objects.filter(period=selected_period).select_related('student__user')
        for account in accounts:
            has_payment = Payment.objects.filter(
                student=account.student, period=selected_period
            ).exists()
            if not has_payment:
                continue
            total_paid = Payment.objects.filter(
                student=account.student, period=selected_period, status='PAID'
            ).aggregate(total=Sum('amount'))['total'] or 0
            total_paid = float(total_paid) if total_paid else 0
            remaining = account.get_display_remaining()
            remaining = float(max(remaining or 0, 0))
            collect = float(account.total_to_collect or 0)
            if remaining > 0:
                is_pending = True
            else:
                is_pending = False
            status_class = 'bg-success' if not is_pending else 'bg-warning text-dark'
            status_text = 'Paid' if not is_pending else 'Due'
            summary.append({
                'student': account.student,
                'total_to_collect': account.total_to_collect,
                'total_paid': total_paid,
                'display_remaining': remaining,
                'remaining_type': account.get_remaining_type(),
                'status_class': status_class,
                'status': status_text,
            })

    return render(request, 'core/payment_summary.html', {
        'period_form': period_form,
        'selected_period': selected_period,
        'summary': summary,
    })


def send_payment_receipt_email(payment):
    if payment.status != 'PAID':
        logger.info(
            "Payment receipt email skipped for payment %s: status is %s",
            payment.id, payment.status,
        )
        return

    student = payment.student
    email = (student.email or '').strip()
    if not email:
        email = getattr(student.user, 'email', '').strip()
    if not email:
        logger.info(
            "Payment receipt email skipped for payment %s: student %s has no email",
            payment.id, student.id,
        )
        return

    display_name = student.get_display_name()
    billing_month = payment.period.name if payment.period_id else 'N/A'

    account, _ = StudentPeriodAccount.objects.get_or_create(
        student=payment.student,
        period=payment.period,
    )
    total_fee = account.total_to_collect or 0
    total_paid = account.get_total_paid() or 0
    remaining = account.get_display_remaining() or 0
    payment_status = 'Paid' if payment.status == 'PAID' else payment.status
    if remaining <= 0:
        status_label = 'PAID IN FULL'
    else:
        status_label = 'DUE'

    payment_date = payment.created_at.strftime('%d %B %Y') if payment.created_at else 'N/A'

    room_no = student.room_no or 'N/A'
    hostel_id = student.hostel_id or 'N/A'
    method = payment.get_method_display() if hasattr(payment, 'get_method_display') else payment.method
    txn_id = payment.txn_id or 'N/A'

    if remaining > 0:
        status_message = (
            f"An amount of ₹{remaining} remains outstanding for this billing period."
        )
    else:
        status_message = (
            "Your mess fee for this billing period has been paid in full. "
            "No outstanding amount remains."
        )

    subject = f"Mess Payment Receipt | {billing_month}"
    plain_body = (
        f"Dear {display_name},\n\n"
        f"Your payment has been successfully recorded. Please find your payment receipt below:\n\n"
        f"PAYMENT DETAILS\n"
        f"Student Name: {display_name}\n"
        f"Hostel ID: {hostel_id}\n"
        f"Room Number: {room_no}\n"
        f"Billing Period: {billing_month}\n"
        f"Payment Date: {payment_date}\n"
        f"Payment Method: {method}\n"
        f"Transaction ID: {txn_id}\n\n"
        f"ACCOUNT SUMMARY\n"
        f"Total Mess Fee: ₹{total_fee}\n"
        f"This Payment: ₹{payment.amount}\n"
        f"Total Amount Paid: ₹{total_paid}\n"
        f"Remaining Due: ₹{remaining}\n"
        f"Payment Status: {status_label}\n\n"
        f"{status_message}\n\n"
        f"Regards,\n"
        f"Mess Committee\n"
        f"APJ Abdul Kalam Boys Hostel\n"
        f"https://hostelmess.in\n"
    )
    html_body = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Payment Receipt</title>
</head>
<body style="margin:0; padding:0; background-color:#f4f4f4; font-family:Arial, Helvetica, sans-serif;">
    <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="background-color:#f4f4f4;">
        <tr>
            <td align="center" style="padding:20px 10px;">
                <table role="presentation" width="600" cellpadding="0" cellspacing="0" border="0" style="background-color:#ffffff; border-radius:8px; overflow:hidden; max-width:600px;">
                    <tr>
                        <td style="background-color:#1a5276; padding:30px 20px; text-align:center;">
                            <h1 style="color:#ffffff; margin:0; font-size:22px; font-weight:bold; line-height:1.3;">APJ Abdul Kalam Boys Hostel</h1>
                            <p style="color:#d6eaf8; margin:8px 0 0 0; font-size:14px; font-weight:bold;">MESS COMMITTEE</p>
                            <p style="color:#aed6f1; margin:4px 0 0 0; font-size:12px;">Hostel Mess Management System</p>
                        </td>
                    </tr>
                    <tr>
                        <td style="padding:30px 20px;">
                            <p style="color:#333333; font-size:16px; margin:0 0 16px 0; line-height:1.5;">Dear {display_name},</p>
                            <p style="color:#555555; font-size:14px; line-height:1.6; margin:0 0 24px 0;">
                                Your payment has been successfully recorded. Please find your payment receipt below:
                            </p>
                            <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="margin:0 0 24px 0;">
                                <tr>
                                    <td style="background-color:#eaf2f8; border-radius:6px; padding:0; overflow:hidden;">
                                        <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0">
                                            <tr>
                                                <td colspan="2" style="background-color:#1a5276; padding:12px 20px; color:#ffffff; font-size:14px; font-weight:bold;">PAYMENT DETAILS</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #d6eaf8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px; width:40%;">Student Name</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold;">{display_name}</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #d6eaf8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Hostel ID</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold;">{hostel_id}</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #d6eaf8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Room Number</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold;">{room_no}</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #d6eaf8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Billing Period</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold;">{billing_month}</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #d6eaf8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Payment Date</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold;">{payment_date}</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #d6eaf8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Payment Method</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold;">{method}</td>
                                            </tr>
                                            <tr>
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Transaction ID</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold;">{txn_id}</td>
                                            </tr>
                                        </table>
                                    </td>
                                </tr>
                            </table>
                            <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="margin:0 0 24px 0;">
                                <tr>
                                    <td style="background-color:#fdfefe; border:1px solid #d5dbdb; border-radius:6px; padding:0; overflow:hidden;">
                                        <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0">
                                            <tr>
                                                <td colspan="2" style="background-color:#1a5276; padding:12px 20px; color:#ffffff; font-size:14px; font-weight:bold;">ACCOUNT SUMMARY</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #e5e8e8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Total Mess Fee</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold; text-align:right;">₹{total_fee}</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #e5e8e8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">This Payment</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold; text-align:right;">₹{payment.amount}</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #e5e8e8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Total Amount Paid</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold; text-align:right;">₹{total_paid}</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #e5e8e8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Remaining Due</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold; text-align:right;">₹{remaining}</td>
                                            </tr>
                                            <tr>
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Payment Status</td>
                                                <td style="padding:10px 20px; font-size:14px; font-weight:bold; text-align:right; color:{'#27ae60' if remaining <= 0 else '#c0392b'};">{status_label}</td>
                                            </tr>
                                        </table>
                                    </td>
                                </tr>
                            </table>
                            <p style="color:#555555; font-size:14px; line-height:1.6; margin:0 0 24px 0;">
                                {status_message}
                            </p>
                        </td>
                    </tr>
                    <tr>
                        <td style="background-color:#f4f4f4; padding:20px; text-align:center; border-top:1px solid #dddddd;">
                            <p style="color:#555555; font-size:12px; margin:0 0 8px 0; line-height:1.5;">
                                Regards,<br>
                                Mess Committee<br>
                                APJ Abdul Kalam Boys Hostel
                            </p>
                            <p style="color:#777777; font-size:12px; margin:0;">
                                Authorized Representative / Mess Committee<br>
                                <a href="https://hostelmess.in" style="color:#1a5276; text-decoration:none;">https://hostelmess.in</a>
                            </p>
                        </td>
                    </tr>
                </table>
            </td>
        </tr>
    </table>
</body>
</html>"""

    try:
        sent = send_email_via_brevo(
            recipient_email=email,
            subject=subject,
            html_content=html_body,
        )
        if sent:
            logger.info(
                "Payment receipt email sent for payment %s to %s",
                payment.id, email,
            )
        else:
            logger.warning(
                "Payment receipt email reported not sent for payment %s to %s",
                payment.id, email,
            )
        return sent
    except Exception:
        logger.exception(
            "Failed to send payment receipt email for payment %s to %s",
            payment.id, email,
        )
        return False


def send_due_payment_email(account):
    remaining = account.get_display_remaining()
    if remaining <= 0:
        logger.info(
            "Due payment email skipped for account %s: remaining is %s",
            account.id, remaining,
        )
        return

    student = account.student
    email = (student.email or '').strip()
    if not email:
        email = getattr(student.user, 'email', '').strip()
    if not email:
        logger.info(
            "Due payment email skipped for account %s: student %s has no email",
            account.id, student.id,
        )
        return

    display_name = student.get_display_name()
    billing_month = account.period.name if account.period_id else 'N/A'
    total_fee = account.total_to_collect or 0
    total_paid = account.get_total_paid() or 0
    remaining_due = remaining or 0

    room_no = student.room_no or 'N/A'
    hostel_id = student.hostel_id or 'N/A'

    subject = f"Mess Fee Due | {billing_month} | ₹{remaining_due} Pending"
    plain_body = (
        f"Dear {display_name},\n\n"
        f"Your mess account currently has an outstanding amount for {billing_month}.\n\n"
        f"PAYMENT SUMMARY\n"
        f"Total Mess Fee: ₹{total_fee}\n"
        f"Total Amount Paid: ₹{total_paid}\n"
        f"Remaining Due: ₹{remaining_due}\n"
        f"Payment Status: DUE\n\n"
        f"PAYMENT REMINDER\n"
        f"An amount of ₹{remaining_due} is still pending for {billing_month}. "
        f"Please clear the pending amount at the earliest to avoid any inconvenience.\n\n"
        f"Student details:\n"
        f"Student Name: {display_name}\n"
        f"Hostel ID: {hostel_id}\n"
        f"Room Number: {room_no}\n"
        f"Billing Period: {billing_month}\n\n"
        f"Regards,\n"
        f"Mess Committee\n"
        f"APJ Abdul Kalam Boys Hostel\n"
        f"Authorized Representative / Mess Committee\n"
        f"https://hostelmess.in\n"
    )
    html_body = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Mess Fee Due</title>
</head>
<body style="margin:0; padding:0; background-color:#f4f4f4; font-family:Arial, Helvetica, sans-serif;">
    <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="background-color:#f4f4f4;">
        <tr>
            <td align="center" style="padding:20px 10px;">
                <table role="presentation" width="600" cellpadding="0" cellspacing="0" border="0" style="background-color:#ffffff; border-radius:8px; overflow:hidden; max-width:600px;">
                    <tr>
                        <td style="background-color:#1a5276; padding:30px 20px; text-align:center;">
                            <h1 style="color:#ffffff; margin:0; font-size:22px; font-weight:bold; line-height:1.3;">APJ Abdul Kalam Boys Hostel</h1>
                            <p style="color:#d6eaf8; margin:8px 0 0 0; font-size:14px; font-weight:bold;">MESS COMMITTEE</p>
                            <p style="color:#aed6f1; margin:4px 0 0 0; font-size:12px;">Hostel Mess Management System</p>
                        </td>
                    </tr>
                    <tr>
                        <td style="padding:30px 20px;">
                            <p style="color:#333333; font-size:16px; margin:0 0 16px 0; line-height:1.5;">Dear {display_name},</p>
                            <p style="color:#555555; font-size:14px; line-height:1.6; margin:0 0 24px 0;">
                                Your mess account currently has an outstanding amount for <strong>{billing_month}</strong>.
                            </p>
                            <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="margin:0 0 24px 0;">
                                <tr>
                                    <td style="background-color:#fdfefe; border:1px solid #d5dbdb; border-radius:6px; padding:0; overflow:hidden;">
                                        <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0">
                                            <tr>
                                                <td colspan="2" style="background-color:#1a5276; padding:12px 20px; color:#ffffff; font-size:14px; font-weight:bold;">PAYMENT SUMMARY</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #e5e8e8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Total Mess Fee</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold; text-align:right;">₹{total_fee}</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #e5e8e8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Total Amount Paid</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold; text-align:right;">₹{total_paid}</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #e5e8e8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Remaining Due</td>
                                                <td style="padding:10px 20px; color:#c0392b; font-size:14px; font-weight:bold; text-align:right;">₹{remaining_due}</td>
                                            </tr>
                                            <tr>
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Payment Status</td>
                                                <td style="padding:10px 20px; color:#c0392b; font-size:14px; font-weight:bold; text-align:right;">DUE</td>
                                            </tr>
                                        </table>
                                    </td>
                                </tr>
                            </table>
                            <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="margin:0 0 24px 0;">
                                <tr>
                                    <td style="background-color:#fdfefe; border:1px solid #d5dbdb; border-radius:6px; padding:0; overflow:hidden;">
                                        <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0">
                                            <tr>
                                                <td colspan="2" style="background-color:#1a5276; padding:12px 20px; color:#ffffff; font-size:14px; font-weight:bold;">PAYMENT REMINDER</td>
                                            </tr>
                                            <tr>
                                                <td style="padding:16px 20px; color:#c0392b; font-size:14px; line-height:1.6;">
                                                    <strong>An amount of ₹{remaining_due} is still pending for {billing_month}.</strong><br>
                                                    Please clear the pending amount at the earliest to avoid any inconvenience.
                                                </td>
                                            </tr>
                                        </table>
                                    </td>
                                </tr>
                            </table>
                            <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="margin:0 0 24px 0;">
                                <tr>
                                    <td style="background-color:#eaf2f8; border-radius:6px; padding:0; overflow:hidden;">
                                        <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0">
                                            <tr>
                                                <td colspan="2" style="background-color:#1a5276; padding:12px 20px; color:#ffffff; font-size:14px; font-weight:bold;">Student Details</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #d6eaf8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px; width:40%;">Student Name</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold;">{display_name}</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #d6eaf8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Hostel ID</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold;">{hostel_id}</td>
                                            </tr>
                                            <tr style="border-bottom:1px solid #d6eaf8;">
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Room Number</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold;">{room_no}</td>
                                            </tr>
                                            <tr>
                                                <td style="padding:10px 20px; color:#555555; font-size:14px;">Billing Period</td>
                                                <td style="padding:10px 20px; color:#333333; font-size:14px; font-weight:bold;">{billing_month}</td>
                                            </tr>
                                        </table>
                                    </td>
                                </tr>
                            </table>
                        </td>
                    </tr>
                    <tr>
                        <td style="background-color:#f4f4f4; padding:20px; text-align:center; border-top:1px solid #dddddd;">
                            <p style="color:#555555; font-size:12px; margin:0 0 8px 0; line-height:1.5;">
                                Regards,<br>
                                Mess Committee<br>
                                APJ Abdul Kalam Boys Hostel
                            </p>
                            <p style="color:#777777; font-size:12px; margin:0;">
                                Authorized Representative / Mess Committee<br>
                                <a href="https://hostelmess.in" style="color:#1a5276; text-decoration:none;">https://hostelmess.in</a>
                            </p>
                        </td>
                    </tr>
                </table>
            </td>
        </tr>
    </table>
</body>
</html>"""

    from_email = getattr(settings, 'DEFAULT_FROM_EMAIL', '') or getattr(settings, 'EMAIL_HOST_USER', '')
    try:
        connection = get_connection(timeout=getattr(settings, 'EMAIL_TIMEOUT', 30))
        sent = send_mail(
            subject,
            plain_body,
            from_email,
            [email],
            connection=connection,
            html_message=html_body,
        )
        if sent:
            logger.info(
                "Due payment email sent for account %s to %s",
                account.id, email,
            )
        else:
            logger.warning(
                "Due payment email reported 0 sent for account %s to %s",
                account.id, email,
            )
    except Exception:
        logger.exception(
            "Failed to send due payment email for account %s to %s",
            account.id, email,
        )


@frontend_management_restricted
def payment_add(request):
    if request.method == 'POST':
        form = PaymentForm(request.POST)
        if form.is_valid():
            try:
                payment = form.save(commit=False)
                if payment.payment_mode == 'custom_full' and payment.period_id and payment.student_id:
                    account, created = StudentPeriodAccount.objects.get_or_create(
                        student=payment.student,
                        period=payment.period,
                    )
                    account.total_to_collect = payment.amount or 0
                    account.is_manual_remaining = False
                    account.manual_remaining = None
                    account.save()
                payment.status = 'PAID'
                payment.save()
            except Exception as e:
                messages.error(request, f'Could not save payment: {e}')
                return render(request, 'core/payment_form.html', {'form': form})

            try:
                send_payment_receipt_email(payment)
            except Exception:
                logger.exception(
                    "Failed to send payment receipt email for payment %s",
                    payment.id,
                )
                messages.warning(request, 'Payment recorded, but receipt email could not be sent.')

            messages.success(request, 'Payment recorded.')
            return redirect('payments_list')
    else:
        form = PaymentForm()
    return render(request, 'core/payment_form.html', {'form': form})


@login_required
@staff_permission_required('can_export_payment_excel')
def payments_summary_export_excel(request):
    return HttpResponse("summary export", content_type='text/plain')


@frontend_management_restricted
def payment_due_edit(request, student_id, period_id):
    account = get_object_or_404(StudentPeriodAccount, student_id=student_id, period_id=period_id)
    form = StudentPeriodAccountForm(request.POST or None, instance=account)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, 'Monthly fee updated.')
        return redirect('payment_summary')
    return render(request, 'core/payment_due_form.html', {'form': form, 'account': account, 'student': account.student, 'period': account.period})


@login_required
@viewer_or_admin_required
def payment_history(request, student_id, period_id):
    account = get_object_or_404(StudentPeriodAccount, student_id=student_id, period_id=period_id)
    qs = Payment.objects.filter(student=account.student, period=account.period).order_by('-created_at')
    ctx = _paginate(request, qs, PER_PAGE)
    ctx['account'] = account
    ctx['student'] = account.student
    return render(request, 'core/payment_history.html', ctx)


@login_required
@viewer_or_admin_required
def due_list(request):
    period_form = PeriodFilterForm(request.GET)
    selected_period = None
    due_list = []
    total_students_with_due = 0
    total_outstanding_due = 0

    if period_form.is_valid():
        selected_period = period_form.cleaned_data.get('period')

    if selected_period:
        accounts = StudentPeriodAccount.objects.filter(period=selected_period).select_related('student__user')
        for account in accounts:
            total_paid = account.get_total_paid()
            remaining = account.get_display_remaining()
            remaining = max(remaining or 0, 0)

            if remaining > 0:
                total_students_with_due += 1
                total_outstanding_due += remaining
                due_list.append({
                    'student': account.student,
                    'total_to_collect': account.total_to_collect,
                    'total_paid': total_paid,
                    'remaining_due': remaining,
                    'remaining_type': account.get_remaining_type(),
                })

    return render(request, 'core/due_list.html', {
        'period_form': period_form,
        'selected_period': selected_period,
        'due_list': due_list,
        'total_students_with_due': total_students_with_due,
        'total_outstanding_due': total_outstanding_due,
    })


def student_search_api(request):
    q = request.GET.get('q', '')
    students = Student.objects.filter(
        Q(hostel_id__icontains=q) | Q(student_name__icontains=q) | Q(user__first_name__icontains=q) | Q(user__last_name__icontains=q)
    )[:10]
    data = [
        {'id': s.id, 'hostel_id': s.hostel_id, 'name': s.student_name or s.user.get_full_name() or s.user.username}
        for s in students
    ]
    return JsonResponse({'results': data})


# Purchases
@login_required
@viewer_or_admin_required
def purchases_list(request):
    qs = Purchase.objects.all().order_by('-bill_date')
    ctx = _paginate(request, qs, PER_PAGE)
    ctx['purchases'] = ctx['page_obj'].object_list
    return render(request, 'core/purchases_list.html', ctx)


@frontend_management_restricted
def purchase_add(request):
    if request.method == 'POST':
        form = PurchaseForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                form.save()
                messages.success(request, 'Purchase recorded.')
                return redirect('purchases_list')
            except Exception as e:
                messages.error(request, f'Could not save purchase: {e}')
    else:
        form = PurchaseForm()
    return render(request, 'core/purchase_form.html', {'form': form})


@frontend_management_restricted
def purchase_edit(request, pk):
    purchase = get_object_or_404(Purchase, pk=pk)
    if request.method == 'POST':
        form = PurchaseForm(request.POST, request.FILES, instance=purchase)
        if form.is_valid():
            form.save()
            messages.success(request, 'Purchase updated.')
            return redirect('purchases_list')
    else:
        form = PurchaseForm(instance=purchase)
    return render(request, 'core/purchase_form.html', {'form': form, 'purchase': purchase})


@login_required
@staff_or_admin_required
def purchases_export_excel(request):
    purchases = list(
        Purchase.objects
        .select_related('supplier', 'item', 'period')
        .all()
        .order_by('-bill_date', '-id')
    )
    generated_at = timezone.now()
    report_date = timezone.localtime(generated_at).strftime('%Y-%m-%d')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'Mess_Purchase_Report_{report_date}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Purchases'
    worksheet.sheet_view.showGridLines = False

    column_count = 11
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
    title_cell = worksheet.cell(
        row=1,
        column=1,
        value='HOSTEL MESS - PURCHASE REPORT',
    )
    title_cell.font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill('solid', fgColor='198754')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    worksheet.row_dimensions[1].height = 28

    worksheet.cell(row=2, column=1, value='Report generated:')
    worksheet.cell(row=2, column=1).font = Font(bold=True, color='404040')
    generated_cell = worksheet.cell(
        row=2,
        column=2,
        value=timezone.localtime(generated_at).replace(tzinfo=None),
    )
    generated_cell.number_format = 'dd mmm yyyy hh:mm:ss'
    generated_cell.alignment = Alignment(horizontal='left')
    worksheet.merge_cells(
        start_row=2,
        start_column=2,
        end_row=2,
        end_column=column_count,
    )

    headers = [
        'Purchase ID',
        'Supplier',
        'Bill No',
        'Bill Date',
        'Accounting Period',
        'Item',
        'Quantity',
        'Unit',
        'Rate',
        'Amount',
        'Bill File',
    ]
    header_row = 4
    header_fill = PatternFill('solid', fgColor='198754')
    header_font = Font(bold=True, color='FFFFFF')
    thin_gray = Side(style='thin', color='D9E2E3')
    cell_border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )

    for column_number, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=header_row, column=column_number, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = cell_border
    worksheet.row_dimensions[header_row].height = 30

    total_amount = Decimal('0.00')
    data_start_row = header_row + 1
    for row_number, purchase in enumerate(purchases, start=data_start_row):
        values = [
            purchase.id,
            purchase.supplier.name if purchase.supplier_id else '',
            purchase.bill_no or '',
            purchase.bill_date,
            purchase.period.name if purchase.period_id else '',
            purchase.item.name if purchase.item_id else '',
            purchase.qty,
            purchase.unit or '',
            purchase.rate,
            purchase.amount,
            purchase.bill_file.name if purchase.bill_file.name else '',
        ]
        total_amount += purchase.amount or Decimal('0.00')

        for column_number, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column_number, value=value)
            cell.border = cell_border
            cell.alignment = Alignment(
                horizontal='right' if column_number in (1, 4, 7, 9, 10) else 'left',
                vertical='center',
                wrap_text=column_number in (2, 5, 6, 11),
            )
        worksheet.cell(row=row_number, column=4).number_format = 'dd-mmm-yyyy'
        worksheet.cell(row=row_number, column=7).number_format = '#,##0.00'
        worksheet.cell(row=row_number, column=9).number_format = '₹ #,##0.00'
        worksheet.cell(row=row_number, column=10).number_format = '₹ #,##0.00'

    data_end_row = data_start_row + len(purchases) - 1
    filter_end_row = max(data_end_row, header_row)
    worksheet.auto_filter.ref = f'A{header_row}:K{filter_end_row}'
    worksheet.freeze_panes = 'A5'

    total_row = data_start_row + len(purchases) + 1
    total_label = worksheet.cell(row=total_row, column=1, value='Total Purchase Amount')
    total_value = worksheet.cell(row=total_row, column=11, value=total_amount)
    total_fill = PatternFill('solid', fgColor='E8F5E9')
    for column_number in range(1, column_count + 1):
        cell = worksheet.cell(row=total_row, column=column_number)
        cell.fill = total_fill
        cell.border = cell_border
        cell.font = Font(bold=True, color='1B5E20')
    total_label.alignment = Alignment(horizontal='right')
    total_value.alignment = Alignment(horizontal='right')
    total_value.number_format = '₹ #,##0.00'
    worksheet.row_dimensions[total_row].height = 22

    column_widths = [14, 22, 16, 15, 22, 22, 13, 12, 15, 16, 24]
    for column_number, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[get_column_letter(column_number)].width = width

    worksheet.print_title_rows = '1:4'
    worksheet.page_setup.orientation = 'landscape'
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.page_margins.top = 0.5
    worksheet.page_margins.bottom = 0.5

    workbook.save(response)
    return response


# Expenses
@login_required
@viewer_or_admin_required
def expenses_list(request):
    labour_qs = LabourPayment.objects.select_related('labour').all().order_by('-month')
    other_qs = OtherExpense.objects.all().order_by('-month')

    labour_ctx = _paginate(request, labour_qs, PER_PAGE, page_param='labour_page')
    other_ctx = _paginate(request, other_qs, PER_PAGE, page_param='other_page')

    ctx = {}
    ctx['labour_expenses'] = labour_ctx['page_obj'].object_list
    ctx['labour_page_obj'] = labour_ctx['page_obj']
    ctx['labour_paginator'] = labour_ctx['page_obj'].paginator
    ctx['labour_is_paginated'] = labour_ctx['is_paginated']
    ctx['labour_query_string'] = labour_ctx['query_string']
    ctx['labour_page_range'] = labour_ctx['page_range_list']

    ctx['other_expenses'] = other_ctx['page_obj'].object_list
    ctx['other_page_obj'] = other_ctx['page_obj']
    ctx['other_paginator'] = other_ctx['page_obj'].paginator
    ctx['other_is_paginated'] = other_ctx['is_paginated']
    ctx['other_query_string'] = other_ctx['query_string']
    ctx['other_page_range'] = other_ctx['page_range_list']

    return render(request, 'core/expenses_list.html', ctx)


@login_required
@staff_permission_required('can_export_expense_excel')
def expenses_export_excel(request):
    return HttpResponse("export expenses", content_type='text/plain')


# Settlements
@login_required
@viewer_or_admin_required
def settlement_list(request):
    qs = MonthlySettlement.objects.all().order_by('-month')
    ctx = _paginate(request, qs, PER_PAGE)
    ctx['settlements'] = ctx['page_obj'].object_list
    ctx['periods'] = AccountingPeriod.objects.order_by('start_date')
    return render(request, 'core/settlement_list.html', ctx)


@login_required
@staff_permission_required('can_manage_settlements')
def settlements_export_excel(request):
    return HttpResponse("export settlements", content_type='text/plain')


@frontend_management_restricted
def auto_settlement(request, year, month, closed_days):
    period = get_object_or_404(AccountingPeriod, start_date__year=year, start_date__month=month)
    # Simplified settlement logic
    students = Student.objects.filter(is_active=True)
    total_fee = 0
    total_expense = 0
    total_paid = 0
    for student in students:
        account, created = StudentPeriodAccount.objects.get_or_create(
            student=student, period=period,
            defaults={'total_to_collect': 0}
        )
        payments = Payment.objects.filter(student=student, period=period, status='PAID')
        paid = payments.aggregate(total=Sum('amount'))['total'] or 0
        total_paid += paid

    purchases = Purchase.objects.filter(period=period)
    total_expense = purchases.aggregate(total=Sum('amount'))['total'] or 0

    labour_payments = LabourPayment.objects.filter(period=period, status='PAID')
    total_expense += labour_payments.aggregate(total=Sum('amount'))['total'] or 0

    other_expenses = OtherExpense.objects.filter(period=period)
    total_expense += other_expenses.aggregate(total=Sum('amount'))['total'] or 0

    per_student = StudentPeriodAccount.objects.filter(period=period).aggregate(
        total=Sum('total_to_collect'))['total'] or 0
    total_fee = per_student

    surplus_deficit = total_fee - total_expense
    num_students = students.count()
    per_student_adjustment = surplus_deficit / num_students if num_students else 0

    settlement, created = MonthlySettlement.objects.get_or_create(
        month=period.start_date.replace(day=1),
        defaults={'accounting_period': period}
    )
    settlement.total_fee = total_fee
    settlement.total_expense = total_expense
    settlement.surplus_deficit = surplus_deficit
    settlement.per_student_adjustment = per_student_adjustment
    settlement.opening_balance = 0
    settlement.closing_balance = surplus_deficit
    settlement.finalized = True
    settlement.save()

    messages.success(request, f'Settlement done for {period.name}.')
    return redirect('settlement_list')


@frontend_management_restricted
def period_settlement(request, period_id):
    period = get_object_or_404(AccountingPeriod, pk=period_id)
    # Reuse auto_settlement logic with 0 closed_days placeholder
    return auto_settlement(request, period.start_date.year, period.start_date.month, 0)


# Reports
@login_required
@staff_permission_required('can_view_reports')
def purchases_report(request):
    qs = Purchase.objects.all().order_by('-bill_date')
    ctx = _paginate(request, qs, PER_PAGE)
    return render(request, 'core/purchases_report.html', ctx)


@login_required
@viewer_or_admin_required
def day_wise_purchase_report(request):
    form = DayWisePurchaseReportForm(request.GET or None)
    purchases = None
    total_amount = 0
    total_entries = 0
    filter_applied = False
    selected_date = None
    date_range = None

    if request.GET and form.is_valid():
        single_date = form.cleaned_data.get('date')
        from_date = form.cleaned_data.get('from_date')
        to_date = form.cleaned_data.get('to_date')

        qs = Purchase.objects.select_related('supplier', 'item', 'period').all()

        if single_date:
            qs = qs.filter(bill_date=single_date)
            selected_date = single_date
            filter_applied = True
        if from_date and to_date:
            qs = qs.filter(bill_date__range=[from_date, to_date])
            date_range = (from_date, to_date)
            filter_applied = True
        if from_date and not to_date:
            qs = qs.filter(bill_date__gte=from_date)
            selected_date = f"From {from_date}"
            filter_applied = True
        if to_date and not from_date:
            qs = qs.filter(bill_date__lte=to_date)
            selected_date = f"Up to {to_date}"
            filter_applied = True

        if filter_applied:
            purchases = qs.order_by('-bill_date')
            total_entries = purchases.count()
            total_amount = qs.aggregate(total=Sum('amount'))['total'] or 0

    return render(request, 'core/day_wise_purchase_report.html', {
        'form': form,
        'purchases': purchases,
        'total_amount': total_amount,
        'total_entries': total_entries,
        'filter_applied': filter_applied,
        'selected_date': selected_date,
        'date_range': date_range,
    })


@login_required
@staff_permission_required('can_view_reports')
def master_report(request):
    return render(request, 'core/master_report.html', {})


@login_required
@staff_or_admin_required
def master_export_excel(request):
    generated_at = timezone.now()
    report_date = timezone.localtime(generated_at).strftime('%Y-%m-%d')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'Hostel_Mess_Master_Report_{report_date}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    workbook = Workbook()

    sheets = ['Summary', 'Payments', 'Purchases', 'Expenses', 'Due List', 'MonthlySettlement']
    for i, sheet_name in enumerate(sheets):
        if i == 0:
            worksheet = workbook.active
            worksheet.title = sheet_name
        else:
            worksheet = workbook.create_sheet(title=sheet_name)
        worksheet.cell(row=1, column=1, value=sheet_name).font = Font(bold=True, size=14)

    # Summary sheet with real data
    ws = workbook['Summary']
    ws.cell(row=1, column=1, value='HOSTEL MESS - MASTER REPORT').font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
    ws.cell(row=1, column=1).fill = PatternFill('solid', fgColor='198754')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    ws.row_dimensions[1].height = 28

    ws.cell(row=2, column=1, value='Report generated:').font = Font(bold=True, color='404040')
    generated_cell = ws.cell(row=2, column=2, value=timezone.localtime(generated_at).replace(tzinfo=None))
    generated_cell.number_format = 'dd mmm yyyy hh:mm:ss'
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=4)

    # Calculations
    total_students = Student.objects.filter(is_active=True).count()
    total_fee_collectable = StudentPeriodAccount.objects.aggregate(total=Sum('total_to_collect'))['total'] or Decimal('0.00')
    total_fee_paid = Payment.objects.filter(status='PAID').aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
    total_purchases = Purchase.objects.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
    total_labour = LabourPayment.objects.filter(status='PAID').aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
    total_other_expense = OtherExpense.objects.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
    total_expenses = total_purchases + total_labour + total_other_expense

    total_due = Decimal('0.00')
    for account in StudentPeriodAccount.objects.all():
        remaining = account.get_display_remaining()
        if remaining:
            total_due += max(remaining, Decimal('0.00'))

    surplus_deficit = total_fee_paid - total_expenses

    thin_gray = Side(style='thin', color='D9E2E3')
    cell_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='198754')
    label_font = Font(bold=True, color='404040')
    value_font = Font(color='1F4E79')
    currency_fmt = '₹ #,##0.00'
    number_fmt = '#,##0'

    row = 4
    ws.cell(row=row, column=1, value='Metric').font = header_font
    ws.cell(row=row, column=1).fill = header_fill
    ws.cell(row=row, column=1).border = cell_border
    ws.cell(row=row, column=1).alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=row, column=2, value='Value').font = header_font
    ws.cell(row=row, column=2).fill = header_fill
    ws.cell(row=row, column=2).border = cell_border
    ws.cell(row=row, column=2).alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 30

    summary_data = [
        ('Total Students', total_students, number_fmt, False),
        ('Total Fee Collectable', total_fee_collectable, currency_fmt, True),
        ('Total Fee Paid', total_fee_paid, currency_fmt, True),
        ('Total Due', total_due, currency_fmt, True),
        ('Total Purchases', total_purchases, currency_fmt, True),
        ('Total Labour Expenses', total_labour, currency_fmt, True),
        ('Total Other Expenses', total_other_expense, currency_fmt, True),
        ('Total Expenses', total_expenses, currency_fmt, True),
        ('Surplus / Deficit', surplus_deficit, currency_fmt, True),
    ]

    for label, value, fmt, is_currency in summary_data:
        row += 1
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = label_font
        label_cell.border = cell_border
        label_cell.alignment = Alignment(horizontal='left', vertical='center')
        val_cell = ws.cell(row=row, column=2, value=value)
        val_cell.font = value_font
        val_cell.border = cell_border
        val_cell.alignment = Alignment(horizontal='right', vertical='center')
        val_cell.number_format = fmt

    # Total row highlight
    total_row = row
    for col in range(1, 3):
        cell = ws.cell(row=total_row, column=col)
        cell.fill = PatternFill('solid', fgColor='E8F5E9')
        cell.font = Font(bold=True, color='1B5E20')

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 10

    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f'A4:B{row}'

    # Payments sheet
    ws_pay = workbook['Payments']
    payments = list(
        Payment.objects
        .select_related('student__user', 'period')
        .all()
        .order_by('-created_at', '-id')
    )

    pay_headers = [
        'Payment ID',
        'Student',
        'Student Name',
        'Hostel ID',
        'Room No',
        'Month',
        'Period',
        'Amount',
        'Adjustment Amount',
        'Method',
        'Status',
        'Transaction ID',
        'Payment Mode',
        'Created At',
    ]
    header_row = 1
    header_fill = PatternFill('solid', fgColor='198754')
    header_font = Font(bold=True, color='FFFFFF')
    thin_gray = Side(style='thin', color='D9E2E3')
    cell_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for col_num, header in enumerate(pay_headers, start=1):
        cell = ws_pay.cell(row=header_row, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = cell_border
    ws_pay.row_dimensions[header_row].height = 30

    total_amount = Decimal('0.00')
    total_adjustment = Decimal('0.00')
    data_start_row = header_row + 1
    for row_num, payment in enumerate(payments, start=data_start_row):
        student_name = payment.student.student_name or payment.student.user.get_full_name() or payment.student.user.username
        hostel_id = payment.student.hostel_id or ''
        room_no = payment.student.room_no or ''
        period_name = payment.period.name if payment.period_id else ''
        values = [
            payment.id,
            f"{hostel_id} - {student_name}" if hostel_id else student_name,
            student_name,
            hostel_id,
            room_no,
            payment.month,
            period_name,
            payment.amount,
            payment.adjustment_amount,
            payment.get_method_display(),
            payment.get_status_display(),
            payment.txn_id or '',
            payment.get_payment_mode_display(),
            timezone.localtime(payment.created_at).replace(tzinfo=None) if payment.created_at else None,
        ]
        total_amount += payment.amount or Decimal('0.00')
        total_adjustment += payment.adjustment_amount or Decimal('0.00')

        for col_num, value in enumerate(values, start=1):
            cell = ws_pay.cell(row=row_num, column=col_num, value=value)
            cell.border = cell_border
            cell.alignment = Alignment(
                horizontal='right' if col_num in (1, 6, 8, 9) else 'left',
                vertical='center',
                wrap_text=col_num in (2, 3, 7, 11, 13),
            )
        ws_pay.cell(row=row_num, column=6).number_format = 'dd-mmm-yyyy'
        ws_pay.cell(row=row_num, column=8).number_format = '₹ #,##0.00'
        ws_pay.cell(row=row_num, column=9).number_format = '₹ #,##0.00'
        ws_pay.cell(row=row_num, column=14).number_format = 'dd mmm yyyy hh:mm:ss'

    data_end_row = data_start_row + len(payments) - 1
    filter_end_row = max(data_end_row, header_row)
    ws_pay.auto_filter.ref = f'A{header_row}:N{filter_end_row}'
    ws_pay.freeze_panes = 'A2'

    # Total row
    total_row = data_start_row + len(payments) + 1
    total_label = ws_pay.cell(row=total_row, column=1, value='Totals')
    total_amt_cell = ws_pay.cell(row=total_row, column=8, value=total_amount)
    total_adj_cell = ws_pay.cell(row=total_row, column=9, value=total_adjustment)
    total_fill = PatternFill('solid', fgColor='E8F5E9')
    for col_num in range(1, len(pay_headers) + 1):
        cell = ws_pay.cell(row=total_row, column=col_num)
        cell.fill = total_fill
        cell.border = cell_border
        cell.font = Font(bold=True, color='1B5E20')
    total_label.alignment = Alignment(horizontal='right')
    total_amt_cell.alignment = Alignment(horizontal='right')
    total_adj_cell.alignment = Alignment(horizontal='right')
    total_amt_cell.number_format = '₹ #,##0.00'
    total_adj_cell.number_format = '₹ #,##0.00'
    ws_pay.row_dimensions[total_row].height = 22

    pay_col_widths = [14, 30, 25, 16, 12, 14, 20, 16, 18, 14, 14, 20, 20, 22]
    for col_num, width in enumerate(pay_col_widths, start=1):
        ws_pay.column_dimensions[get_column_letter(col_num)].width = width

    ws_pay.print_title_rows = '1:1'
    ws_pay.page_setup.orientation = 'landscape'
    ws_pay.page_setup.fitToWidth = 1
    ws_pay.page_setup.fitToHeight = 0
    ws_pay.sheet_properties.pageSetUpPr.fitToPage = True
    ws_pay.page_margins.left = 0.25
    ws_pay.page_margins.right = 0.25
    ws_pay.page_margins.top = 0.5
    ws_pay.page_margins.bottom = 0.5

    # Purchases sheet
    ws_pur = workbook['Purchases']
    purchases = list(
        Purchase.objects
        .select_related('supplier', 'item', 'period')
        .all()
        .order_by('-bill_date', '-id')
    )

    pur_headers = [
        'Purchase ID',
        'Supplier',
        'Bill No',
        'Bill Date',
        'Period',
        'Item',
        'Quantity',
        'Unit',
        'Rate',
        'Amount',
        'Bill File',
    ]
    header_row = 1
    header_fill = PatternFill('solid', fgColor='198754')
    header_font = Font(bold=True, color='FFFFFF')
    thin_gray = Side(style='thin', color='D9E2E3')
    cell_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for col_num, header in enumerate(pur_headers, start=1):
        cell = ws_pur.cell(row=header_row, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = cell_border
    ws_pur.row_dimensions[header_row].height = 30

    total_amount = Decimal('0.00')
    data_start_row = header_row + 1
    for row_num, purchase in enumerate(purchases, start=data_start_row):
        values = [
            purchase.id,
            purchase.supplier.name if purchase.supplier_id else '',
            purchase.bill_no or '',
            purchase.bill_date,
            purchase.period.name if purchase.period_id else '',
            purchase.item.name if purchase.item_id else '',
            purchase.qty,
            purchase.unit or '',
            purchase.rate,
            purchase.amount,
            purchase.bill_file.name if purchase.bill_file.name else '',
        ]
        total_amount += purchase.amount or Decimal('0.00')

        for col_num, value in enumerate(values, start=1):
            cell = ws_pur.cell(row=row_num, column=col_num, value=value)
            cell.border = cell_border
            cell.alignment = Alignment(
                horizontal='right' if col_num in (1, 4, 7, 9, 10) else 'left',
                vertical='center',
                wrap_text=col_num in (2, 5, 6, 11),
            )
        ws_pur.cell(row=row_num, column=4).number_format = 'dd-mmm-yyyy'
        ws_pur.cell(row=row_num, column=7).number_format = '#,##0.00'
        ws_pur.cell(row=row_num, column=9).number_format = '₹ #,##0.00'
        ws_pur.cell(row=row_num, column=10).number_format = '₹ #,##0.00'

    data_end_row = data_start_row + len(purchases) - 1
    filter_end_row = max(data_end_row, header_row)
    ws_pur.auto_filter.ref = f'A{header_row}:K{filter_end_row}'
    ws_pur.freeze_panes = 'A2'

    # Total row
    total_row = data_start_row + len(purchases) + 1
    total_label = ws_pur.cell(row=total_row, column=1, value='Total Purchase Amount')
    total_value = ws_pur.cell(row=total_row, column=11, value=total_amount)
    total_fill = PatternFill('solid', fgColor='E8F5E9')
    for col_num in range(1, len(pur_headers) + 1):
        cell = ws_pur.cell(row=total_row, column=col_num)
        cell.fill = total_fill
        cell.border = cell_border
        cell.font = Font(bold=True, color='1B5E20')
    total_label.alignment = Alignment(horizontal='right')
    total_value.alignment = Alignment(horizontal='right')
    total_value.number_format = '₹ #,##0.00'
    ws_pur.row_dimensions[total_row].height = 22

    pur_col_widths = [14, 22, 16, 15, 22, 22, 13, 12, 15, 16, 24]
    for col_num, width in enumerate(pur_col_widths, start=1):
        ws_pur.column_dimensions[get_column_letter(col_num)].width = width

    ws_pur.print_title_rows = '1:1'
    ws_pur.page_setup.orientation = 'landscape'
    ws_pur.page_setup.fitToWidth = 1
    ws_pur.page_setup.fitToHeight = 0
    ws_pur.sheet_properties.pageSetUpPr.fitToPage = True
    ws_pur.page_margins.left = 0.25
    ws_pur.page_margins.right = 0.25
    ws_pur.page_margins.top = 0.5
    ws_pur.page_margins.bottom = 0.5

    # Expenses sheet
    ws_exp = workbook['Expenses']

    labour_payments = list(
        LabourPayment.objects
        .select_related('labour', 'period')
        .filter(status='PAID')
        .all()
        .order_by('-month', '-id')
    )

    other_expenses = list(
        OtherExpense.objects
        .select_related('period')
        .all()
        .order_by('-month', '-id')
    )

    exp_headers = [
        'Source / Type',
        'Labour / Category',
        'Month',
        'Period',
        'Amount',
        'Status / Note',
    ]
    header_row = 1
    header_fill = PatternFill('solid', fgColor='198754')
    header_font = Font(bold=True, color='FFFFFF')
    thin_gray = Side(style='thin', color='D9E2E3')
    cell_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for col_num, header in enumerate(exp_headers, start=1):
        cell = ws_exp.cell(row=header_row, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = cell_border
    ws_exp.row_dimensions[header_row].height = 30

    total_expenses = Decimal('0.00')
    data_start_row = header_row + 1
    row_num = data_start_row

    # Labour Payments
    for lp in labour_payments:
        values = [
            'Labour Payment',
            lp.labour.name if lp.labour_id else '',
            lp.month,
            lp.period.name if lp.period_id else '',
            lp.amount,
            lp.get_status_display(),
        ]
        total_expenses += lp.amount or Decimal('0.00')

        for col_num, value in enumerate(values, start=1):
            cell = ws_exp.cell(row=row_num, column=col_num, value=value)
            cell.border = cell_border
            cell.alignment = Alignment(
                horizontal='right' if col_num == 5 else 'left',
                vertical='center',
                wrap_text=col_num in (2, 6),
            )
        ws_exp.cell(row=row_num, column=3).number_format = 'mmm-yyyy'
        ws_exp.cell(row=row_num, column=5).number_format = '₹ #,##0.00'
        row_num += 1

    # Other Expenses
    for oe in other_expenses:
        values = [
            'Other Expense',
            oe.category or '',
            oe.month,
            oe.period.name if oe.period_id else '',
            oe.amount,
            oe.note or '',
        ]
        total_expenses += oe.amount or Decimal('0.00')

        for col_num, value in enumerate(values, start=1):
            cell = ws_exp.cell(row=row_num, column=col_num, value=value)
            cell.border = cell_border
            cell.alignment = Alignment(
                horizontal='right' if col_num == 5 else 'left',
                vertical='center',
                wrap_text=col_num in (2, 6),
            )
        ws_exp.cell(row=row_num, column=3).number_format = 'mmm-yyyy'
        ws_exp.cell(row=row_num, column=5).number_format = '₹ #,##0.00'
        row_num += 1

    data_end_row = data_start_row + len(labour_payments) + len(other_expenses) - 1
    filter_end_row = max(data_end_row, header_row)
    ws_exp.auto_filter.ref = f'A{header_row}:F{filter_end_row}'
    ws_exp.freeze_panes = 'A2'

    # Total row
    total_row = data_start_row + len(labour_payments) + len(other_expenses) + 1
    total_label = ws_exp.cell(row=total_row, column=1, value='Total Expenses')
    total_value = ws_exp.cell(row=total_row, column=5, value=total_expenses)
    total_fill = PatternFill('solid', fgColor='E8F5E9')
    for col_num in range(1, len(exp_headers) + 1):
        cell = ws_exp.cell(row=total_row, column=col_num)
        cell.fill = total_fill
        cell.border = cell_border
        cell.font = Font(bold=True, color='1B5E20')
    total_label.alignment = Alignment(horizontal='right')
    total_value.alignment = Alignment(horizontal='right')
    total_value.number_format = '₹ #,##0.00'
    ws_exp.row_dimensions[total_row].height = 22

    exp_col_widths = [20, 30, 15, 22, 16, 40]
    for col_num, width in enumerate(exp_col_widths, start=1):
        ws_exp.column_dimensions[get_column_letter(col_num)].width = width

    ws_exp.print_title_rows = '1:1'
    ws_exp.page_setup.orientation = 'landscape'
    ws_exp.page_setup.fitToWidth = 1
    ws_exp.page_setup.fitToHeight = 0
    ws_exp.sheet_properties.pageSetUpPr.fitToPage = True
    ws_exp.page_margins.left = 0.25
    ws_exp.page_margins.right = 0.25
    ws_exp.page_margins.top = 0.5
    ws_exp.page_margins.bottom = 0.5

    # Due List sheet
    ws_due = workbook['Due List']

    accounts = list(
        StudentPeriodAccount.objects
        .select_related('student__user', 'period')
        .all()
        .order_by('period__start_date', 'student__hostel_id')
    )

    due_headers = [
        'Student',
        'Student Name',
        'Hostel ID',
        'Room No',
        'Period',
        'Total Fee',
        'Total Paid',
        'Due',
        'Remaining Type',
        'Status',
    ]
    header_row = 1
    header_fill = PatternFill('solid', fgColor='198754')
    header_font = Font(bold=True, color='FFFFFF')
    thin_gray = Side(style='thin', color='D9E2E3')
    cell_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for col_num, header in enumerate(due_headers, start=1):
        cell = ws_due.cell(row=header_row, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = cell_border
    ws_due.row_dimensions[header_row].height = 30

    total_fee = Decimal('0.00')
    total_paid = Decimal('0.00')
    total_due = Decimal('0.00')
    data_start_row = header_row + 1
    row_num = data_start_row

    for account in accounts:
        total_paid_amount = account.get_total_paid() or Decimal('0.00')
        remaining = account.get_display_remaining()
        remaining = max(remaining or Decimal('0.00'), Decimal('0.00'))
        collect = account.total_to_collect or Decimal('0.00')

        student_name = account.student.student_name or account.student.user.get_full_name() or account.student.user.username
        hostel_id = account.student.hostel_id or ''
        room_no = account.student.room_no or ''
        period_name = account.period.name if account.period_id else ''
        remaining_type = account.get_remaining_type()

        if remaining > 0:
            status = 'Due'
        else:
            status = 'Paid'

        values = [
            f"{hostel_id} - {student_name}" if hostel_id else student_name,
            student_name,
            hostel_id,
            room_no,
            period_name,
            collect,
            total_paid_amount,
            remaining,
            remaining_type,
            status,
        ]
        total_fee += collect
        total_paid += total_paid_amount
        total_due += remaining

        for col_num, value in enumerate(values, start=1):
            cell = ws_due.cell(row=row_num, column=col_num, value=value)
            cell.border = cell_border
            cell.alignment = Alignment(
                horizontal='right' if col_num in (6, 7, 8) else 'left',
                vertical='center',
                wrap_text=col_num in (1, 2, 5, 9, 10),
            )
        ws_due.cell(row=row_num, column=6).number_format = '₹ #,##0.00'
        ws_due.cell(row=row_num, column=7).number_format = '₹ #,##0.00'
        ws_due.cell(row=row_num, column=8).number_format = '₹ #,##0.00'
        row_num += 1

    data_end_row = data_start_row + len(accounts) - 1
    filter_end_row = max(data_end_row, header_row)
    ws_due.auto_filter.ref = f'A{header_row}:J{filter_end_row}'
    ws_due.freeze_panes = 'A2'

    # Total row
    total_row = data_start_row + len(accounts) + 1
    total_label = ws_due.cell(row=total_row, column=1, value='Totals')
    total_fee_cell = ws_due.cell(row=total_row, column=6, value=total_fee)
    total_paid_cell = ws_due.cell(row=total_row, column=7, value=total_paid)
    total_due_cell = ws_due.cell(row=total_row, column=8, value=total_due)
    total_fill = PatternFill('solid', fgColor='E8F5E9')
    for col_num in range(1, len(due_headers) + 1):
        cell = ws_due.cell(row=total_row, column=col_num)
        cell.fill = total_fill
        cell.border = cell_border
        cell.font = Font(bold=True, color='1B5E20')
    total_label.alignment = Alignment(horizontal='right')
    total_fee_cell.alignment = Alignment(horizontal='right')
    total_paid_cell.alignment = Alignment(horizontal='right')
    total_due_cell.alignment = Alignment(horizontal='right')
    total_fee_cell.number_format = '₹ #,##0.00'
    total_paid_cell.number_format = '₹ #,##0.00'
    total_due_cell.number_format = '₹ #,##0.00'
    ws_due.row_dimensions[total_row].height = 22

    due_col_widths = [30, 25, 16, 12, 22, 16, 16, 16, 18, 14]
    for col_num, width in enumerate(due_col_widths, start=1):
        ws_due.column_dimensions[get_column_letter(col_num)].width = width

    ws_due.print_title_rows = '1:1'
    ws_due.page_setup.orientation = 'landscape'
    ws_due.page_setup.fitToWidth = 1
    ws_due.page_setup.fitToHeight = 0
    ws_due.sheet_properties.pageSetUpPr.fitToPage = True
    ws_due.page_margins.left = 0.25
    ws_due.page_margins.right = 0.25
    ws_due.page_margins.top = 0.5
    ws_due.page_margins.bottom = 0.5

    # MonthlySettlement sheet
    ws_set = workbook['MonthlySettlement']

    settlements = list(
        MonthlySettlement.objects
        .select_related('accounting_period')
        .all()
        .order_by('-month', '-id')
    )

    set_headers = [
        'Settlement ID',
        'Month',
        'Accounting Period',
        'Total Fee',
        'Total Expense',
        'Surplus / Deficit',
        'Per Student Adjustment',
        'Opening Balance',
        'Closing Balance',
        'Finalized',
    ]
    header_row = 1
    header_fill = PatternFill('solid', fgColor='198754')
    header_font = Font(bold=True, color='FFFFFF')
    thin_gray = Side(style='thin', color='D9E2E3')
    cell_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for col_num, header in enumerate(set_headers, start=1):
        cell = ws_set.cell(row=header_row, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = cell_border
    ws_set.row_dimensions[header_row].height = 30

    data_start_row = header_row + 1
    for row_num, settlement in enumerate(settlements, start=data_start_row):
        period_name = settlement.accounting_period.name if settlement.accounting_period_id else ''
        values = [
            settlement.id,
            settlement.month,
            period_name,
            settlement.total_fee,
            settlement.total_expense,
            settlement.surplus_deficit,
            settlement.per_student_adjustment,
            settlement.opening_balance,
            settlement.closing_balance,
            'Yes' if settlement.finalized else 'No',
        ]

        for col_num, value in enumerate(values, start=1):
            cell = ws_set.cell(row=row_num, column=col_num, value=value)
            cell.border = cell_border
            cell.alignment = Alignment(
                horizontal='right' if col_num in (1, 4, 5, 6, 7, 8, 9) else 'left',
                vertical='center',
                wrap_text=col_num in (3, 10),
            )
        ws_set.cell(row=row_num, column=2).number_format = 'mmm-yyyy'
        for col_num in (4, 5, 6, 7, 8, 9):
            ws_set.cell(row=row_num, column=col_num).number_format = '₹ #,##0.00'

    data_end_row = data_start_row + len(settlements) - 1
    filter_end_row = max(data_end_row, header_row)
    ws_set.auto_filter.ref = f'A{header_row}:J{filter_end_row}'
    ws_set.freeze_panes = 'A2'

    set_col_widths = [16, 15, 22, 16, 16, 18, 20, 18, 18, 14]
    for col_num, width in enumerate(set_col_widths, start=1):
        ws_set.column_dimensions[get_column_letter(col_num)].width = width

    ws_set.print_title_rows = '1:1'
    ws_set.page_setup.orientation = 'landscape'
    ws_set.page_setup.fitToWidth = 1
    ws_set.page_setup.fitToHeight = 0
    ws_set.sheet_properties.pageSetUpPr.fitToPage = True
    ws_set.page_margins.left = 0.25
    ws_set.page_margins.right = 0.25
    ws_set.page_margins.top = 0.5
    ws_set.page_margins.bottom = 0.5

    workbook.save(response)
    return response


# User Activity (Admin Only)
@frontend_management_restricted
def user_activity(request):
    profiles = UserProfile.objects.select_related('user').order_by('-updated_at')
    ctx = _paginate(request, profiles, PER_PAGE)
    ctx['profiles'] = ctx['page_obj']
    ctx.pop('page_obj', None)
    return render(request, 'core/user_activity.html', ctx)


# User Management (Admin Only)
@frontend_management_restricted
def user_manage(request):
    profiles = UserProfile.objects.select_related('user').all()
    return render(request, 'core/user_manage.html', {'profiles': profiles})


@frontend_management_restricted
def user_change_role(request, user_id):
    profile = get_object_or_404(UserProfile, user_id=user_id)
    if request.method == 'POST':
        role = request.POST.get('role')
        if role in ['admin', 'staff', 'viewer']:
            profile.role = role
            profile.save()
            messages.success(request, f'Role changed to {role}.')
    return redirect('user_manage')


@frontend_management_restricted
def user_toggle_account(request, user_id):
    profile = get_object_or_404(UserProfile, user_id=user_id)
    profile.is_active_user = not profile.is_active_user
    profile.save()
    messages.success(request, 'Account status toggled.')
    return redirect('user_manage')


@frontend_management_restricted
def user_permissions_edit(request, user_id):
    profile = get_object_or_404(UserProfile, user_id=user_id)
    if request.method == 'POST':
        for field in [
            'can_manage_students', 'can_manage_payments', 'can_manage_expenses',
            'can_manage_purchases', 'can_manage_settlements', 'can_view_reports',
            'can_export_collection_excel', 'can_export_expense_excel',
            'can_export_purchase_excel', 'can_export_payment_excel',
        ]:
            setattr(profile, field, request.POST.get(field) == 'on')
        profile.save()
        messages.success(request, 'Permissions updated.')
    return render(request, 'core/user_permissions_edit.html', {'profile': profile})
